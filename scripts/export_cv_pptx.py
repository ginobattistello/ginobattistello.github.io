#!/usr/bin/env python3
"""Create an editable A4 portrait PowerPoint CV from the site's YAML content.

Install scripts/requirements-cv.txt, then run this script from any directory.
Existing files are preserved unless --overwrite is explicitly supplied.
"""
import argparse
from pathlib import Path

import yaml
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
PAGE_W, PAGE_H = 595.276, 841.89
LEFT, TOP, WIDTH, END = 45.354, 42.52, 504.568, 790.0
DATE_W, BODY_X = 68.032, LEFT + 73.701
BODY_W = WIDTH - (BODY_X - LEFT)
TEXT, MUTED, RULE, ORANGE = '111111', '666666', 'E2E2E2', 'F26A1B'
BODY, META, SMALL, SUBHEAD = 9.6, 9.0, 7.2, 10.0
FONT = 'Arial'
PALETTE = {
    'Grants': ('7A2E00', 'FFB36B'), 'Awards': ('FFE4CC', '8A3800'),
    'Organiser': ('444444', 'FFFFFF'), 'Co-organiser': ('444444', 'FFFFFF'),
    'Interview': ('D8D8D8', '333333'), 'Volunteer': ('F0F0F0', '595959'),
    'Talk': ('111111', 'FFFFFF'),
}


def load(name):
    return yaml.safe_load((ROOT / '_data' / (name + '.yml')).read_text())


def font_path(bold=False):
    names = [
        '/System/Library/Fonts/Supplemental/Arial' + (' Bold' if bold else '') + '.ttf',
        'C:/Windows/Fonts/' + ('arialbd.ttf' if bold else 'arial.ttf'),
        '/usr/share/fonts/truetype/dejavu/DejaVuSans' + ('-Bold' if bold else '') + '.ttf',
    ]
    return next((p for p in names if Path(p).is_file()), None)


def measure(text, size, bold=False):
    path = font_path(bold)
    font = ImageFont.truetype(path, round(size * 8)) if path else ImageFont.load_default()
    scale = 8 if path else 1
    return font.getlength(str(text)) / scale


def line_count(text, width, size, bold=False):
    """Conservative measurement leaves room for Office font/line-break differences."""
    total = 0
    for paragraph in str(text).split('\n'):
        line = ''
        for word in paragraph.split():
            candidate = (line + ' ' + word).strip()
            if line and measure(candidate, size, bold) > width * .94:
                total += 1
                line = word
            else:
                line = candidate
            while measure(line, size, bold) > width * .94:
                split = max(1, int(len(line) * width * .90 / measure(line, size, bold)))
                total += 1
                line = line[split:]
        total += 1
    return max(total, 1)


def text_height(text, width, size=BODY, bold=False):
    return line_count(text, width, size, bold) * size * 1.36 + 2


def set_font(run, size, bold=False, color=TEXT):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


class CV:
    def __init__(self, profile):
        self.profile = profile
        self.deck = Presentation()
        self.deck.slide_width, self.deck.slide_height = Pt(PAGE_W), Pt(PAGE_H)
        self.deck.core_properties.title = profile['name'] + ' — editable curriculum vitae'
        self.deck.core_properties.author = profile['name']
        self.section = None
        self.subsection = None
        self.new_page()

    def new_page(self):
        self.slide = self.deck.slides.add_slide(self.deck.slide_layouts[6])
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        self.y = TOP

    def text(self, text, x, y, width, size=BODY, bold=False, color=TEXT,
             height=None, link=None, align=PP_ALIGN.LEFT, emphasize=None):
        height = height or text_height(text, width, size, bold)
        shape = self.slide.shapes.add_textbox(Pt(x), Pt(y), Pt(width), Pt(height))
        shape.name = str(text).split('\n')[0][:80]
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.TOP
        for index, line in enumerate(str(text).split('\n')):
            p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            p.alignment = align
            p.line_spacing = Pt(size * 1.36)
            p.space_before = p.space_after = Pt(0)
            parts = line.split(emphasize) if emphasize and emphasize in line else [line]
            for n, part in enumerate(parts):
                if n:
                    run = p.add_run()
                    run.text = emphasize
                    set_font(run, size, True, color)
                run = p.add_run()
                run.text = part
                set_font(run, size, bold, color)
                if link:
                    run.hyperlink.address = link
        return height

    def rule(self, y, x=LEFT, width=WIDTH):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(width), Pt(.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(RULE)
        shape.line.fill.background()

    def ensure(self, height):
        if height > END - TOP - 55:
            raise ValueError('One entry is too tall for an A4 page; split it in the source data.')
        if self.y + height > END:
            self.new_page()
            if self.section:
                self.heading(self.section + ' (continued)', continuation=True)
            if self.subsection:
                self.subheading(self.subsection + ' (continued)', continuation=True)

    def header(self):
        p = self.profile
        self.text(p['name'], LEFT, self.y, 300, 24)
        self.text(p['role'] + '\n' + p['affiliation'], LEFT, self.y + 36, 300, 10, color=MUTED)
        contact = [p['email'], 'ginobattistello.github.io', 'ORCID 0009-0009-9355-470X',
                   'Google Scholar · GitHub', 'Born ' + str(p['birth_year']) + ' · Nationality: ' + p['nationality']]
        right_height = self.text('\n'.join(contact), LEFT + WIDTH - 190, self.y, 190, 9,
                                 color=MUTED, align=PP_ALIGN.RIGHT)
        self.y += max(36 + text_height(p['role'] + '\n' + p['affiliation'], 300, 10), right_height) + 9
        self.rule(self.y)
        self.y += 16

    def heading(self, title, continuation=False):
        if not continuation:
            if self.y + 65 > END:
                self.new_page()
            self.section, self.subsection = title, None
        square = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(LEFT), Pt(self.y + 4), Pt(4.5), Pt(4.5))
        square.fill.solid()
        square.fill.fore_color.rgb = RGBColor.from_string(ORANGE)
        square.line.fill.background()
        self.text(title.upper(), LEFT + 12, self.y, WIDTH - 12, 8.4, True)
        self.rule(self.y + 16)
        self.y += 24

    def subheading(self, title, divider=False, continuation=False):
        if not continuation:
            self.ensure(50)
            self.subsection = title
        if divider:
            self.rule(self.y + 3)
            self.y += 10
        self.text(title, BODY_X, self.y + 2, BODY_W, SUBHEAD, True, MUTED)
        self.y += text_height(title, BODY_W, SUBHEAD, True) + 5

    def pill_width(self, label):
        return max(34, measure(label.upper(), SMALL, True) + 9)

    def pill(self, label, x, y, width):
        bg, fg = PALETTE.get(label, ('FFFFFF', MUTED))
        shape = self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y + 1), Pt(width), Pt(12))
        shape.adjustments[0] = .12
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(bg)
        if bg == 'FFFFFF':
            shape.line.color.rgb = RGBColor.from_string(RULE)
            shape.line.width = Pt(.5)
        else:
            shape.line.fill.background()
        self.text(label.upper(), x + 2, y + 1, width - 4, SMALL, True, fg, height=12, align=PP_ALIGN.CENTER)

    def entry(self, entry, plain=False, divider=False, extra=None):
        lines = [('title', entry['title'], BODY, not plain, TEXT, entry.get('tag'), entry.get('url'))]
        if entry.get('subtitle'):
            lines.append(('text', entry['subtitle'], BODY, False, TEXT, None, None))
        for note in entry.get('notes', []):
            lines.append(('text', note, META, False, MUTED, None, None))
        meta = ' · '.join(str(entry[k]) for k in ('people', 'place') if entry.get(k))
        if meta:
            lines.append(('text', meta, META, False, MUTED, None, None))
        lines += extra or []
        for certificate in entry.get('certificates', []):
            lines.append(('text', certificate, META, False, MUTED, 'Certificate', None))
        heights = []
        for _, text, size, bold, _, label, _ in lines:
            width = BODY_W - (self.pill_width(label) + 5 if label else 0)
            heights.append(max(14 if label else 0, text_height(text, width, size, bold)))
        date = str(entry.get('date', ''))
        if entry.get('duration'):
            date += '\n' + str(entry['duration'])
        height = max(sum(heights) + 1.5 * (len(lines) - 1), text_height(date, DATE_W, META)) + 6
        self.ensure(height + (11 if divider else 0))
        if divider:
            self.rule(self.y + 2)
            self.y += 11
        first_shape = len(self.slide.shapes)
        self.text(date, LEFT, self.y, DATE_W, META, color=MUTED)
        for (kind, text, size, bold, color, label, link), h in zip(lines, heights):
            offset = self.pill_width(label) + 5 if label else 0
            if label:
                self.pill(label, BODY_X, self.y, offset - 5)
            self.text(text, BODY_X + offset, self.y, BODY_W - offset, size, bold, color,
                      height=h, link=link, emphasize=self.profile.get('author_name') if kind == 'authors' else None)
            self.y += h + 1.5
        self.y += 4.5
        group = self.slide.shapes.add_group_shape(list(self.slide.shapes)[first_shape:])
        group.name = entry['title'][:80]

    def entries(self, entries, plain=False):
        previous = None
        for entry in entries:
            self.entry(entry, plain, bool(previous and entry.get('stage') != previous))
            previous = entry.get('stage')
        self.y += 6

    def skills(self, domains):
        self.new_page()
        self.heading('Research expertise & technical skills')
        category_w = 121.89
        for domain in domains:
            rows = []
            for item in domain['items']:
                detail_height = sum(text_height((d.get('label', '') + ': ' if d.get('label') else '') + d['text'],
                                               WIDTH - category_w - 12, META if d.get('level') == 'tools' else BODY)
                                    for d in item['details']) + 2 * (len(item['details']) - 1)
                category_height = text_height(item['category'], category_w - 12, 9.4, True)
                rows.append(max(detail_height, category_height) + 10)
            self.ensure(min(sum(rows) + 30, 80))
            self.text(domain['domain'], LEFT + 6, self.y + 3, WIDTH - 12, SUBHEAD, True, MUTED)
            self.y += 23
            # Each row is a real, independently editable PowerPoint table.
            for item, height in zip(domain['items'], rows):
                self.ensure(height)
                shape = self.slide.shapes.add_table(1, 2, Pt(LEFT), Pt(self.y), Pt(WIDTH), Pt(height))
                table = shape.table
                table.first_row = False
                table.horz_banding = False
                table.columns[0].width, table.columns[1].width = Pt(category_w), Pt(WIDTH - category_w)
                for index, cell in enumerate(table.rows[0].cells):
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
                    cell.margin_left = cell.margin_right = Pt(6)
                    cell.margin_top = cell.margin_bottom = Pt(4)
                    cell.vertical_anchor = MSO_ANCHOR.TOP
                    cell.text_frame.clear()
                    cell.text_frame.word_wrap = True
                    cell.text_frame.auto_size = MSO_AUTO_SIZE.NONE
                    details = [{'text': item['category']}] if index == 0 else item['details']
                    for n, detail in enumerate(details):
                        p = cell.text_frame.paragraphs[0] if n == 0 else cell.text_frame.add_paragraph()
                        size = 9.4 if index == 0 else META if detail.get('level') == 'tools' else BODY
                        color = MUTED if detail.get('level') == 'tools' else TEXT
                        p.line_spacing = Pt(size * 1.36)
                        p.space_before, p.space_after = Pt(0), Pt(2 if n < len(details) - 1 else 0)
                        if detail.get('label'):
                            run = p.add_run()
                            run.text = detail['label'] + ': '
                            set_font(run, size, True, color)
                        run = p.add_run()
                        run.text = detail['text']
                        set_font(run, size, index == 0, color)
                    props = cell._tc.get_or_add_tcPr()
                    for edge in ('L', 'R', 'T', 'B'):
                        line = OxmlElement('a:ln' + edge)
                        line.set('w', '6350')
                        fill = OxmlElement('a:solidFill' if edge == 'B' else 'a:noFill')
                        if edge == 'B':
                            color = OxmlElement('a:srgbClr')
                            color.set('val', RULE)
                            fill.append(color)
                        line.append(fill)
                        props.append(line)
                self.y += height
            self.y += 5

    def finish(self, output):
        count = len(self.deck.slides)
        for index, slide in enumerate(self.deck.slides, 1):
            self.slide = slide
            self.text(self.profile['name'] + ' · Curriculum Vitae', LEFT, PAGE_H - 27, WIDTH - 50, 8, color='9A9A9A')
            self.text(f'{index} / {count}', LEFT + WIDTH - 45, PAGE_H - 27, 45, 8,
                      color='9A9A9A', align=PP_ALIGN.RIGHT)
            for shape in slide.shapes:
                if shape.left < 0 or shape.top < 0 or shape.left + shape.width > Pt(PAGE_W + .1) or shape.top + shape.height > Pt(PAGE_H + .1):
                    raise ValueError(f'Shape exceeds slide boundary on page {index}: {shape.name}')
        output.parent.mkdir(parents=True, exist_ok=True)
        self.deck.save(output)
        print(f'Created {output} ({count} editable A4 portrait slides).')


def build(output):
    profile, education = load('profile'), load('education')
    cv = CV(profile)
    cv.header()
    for title, entries in [('Research experience', load('experience')), ('Education', education['degrees']),
                           ('Software & tools', load('software')), ('Team responsibilities', load('team'))]:
        if entries:
            cv.heading(title)
            cv.entries(entries)
    cv.heading('Publications')
    for pub in load('publications')['articles']:
        venue = pub['venue'] + (', ' + str(pub['details']) if pub.get('details') else '')
        if pub.get('doi'):
            venue += ' · doi:' + pub['doi']
        cv.entry({'date': pub['year'], 'title': pub['title']}, extra=[
            ('authors', pub['authors'], META, False, MUTED, None, None),
            ('text', venue, META, False, MUTED, None, 'https://doi.org/' + pub['doi'] if pub.get('doi') else None)])
    cv.subheading('Dissertations')
    for pub in load('publications')['theses']:
        cv.entry({'date': pub['year'], 'title': pub['title']}, extra=[
            ('authors', pub['authors'], META, False, MUTED, None, None),
            ('text', pub['venue'], META, False, MUTED, None, None)])
    cv.heading('Talks & posters')
    for talk in load('presentations'):
        extra = [('authors', talk['authors'], META, False, MUTED, None, None)]
        for venue in talk['venues']:
            text = venue['name'] + ', ' + venue['place']
            if str(venue['year']) != str(talk['year']):
                text += ' (' + str(venue['year']) + ')'
            if venue.get('note'):
                text += ' · ' + venue['note']
            extra.append(('text', text, META, False, TEXT, venue['type'], None))
        cv.entry({'date': talk['year'], 'title': talk['title']}, extra=extra)
    cv.heading('Training')
    cv.entries(education['training'])
    cv.heading('Teaching & supervision')
    cv.subheading('Teaching')
    cv.entries(load('teaching')['teaching'])
    cv.subheading('Student supervision', divider=True)
    cv.entries(load('teaching')['supervision'], plain=True)
    cv.heading('Awards & grants')
    cv.entries(education['awards'])
    cv.heading('Service & outreach')
    for index, group in enumerate(load('outreach')):
        cv.subheading(group['group'] + (' · ' + group['intro'] if group.get('intro') else ''), divider=index > 0)
        cv.entries(group['items'], plain=True)
    cv.heading('Languages')
    cv.text(' · '.join(education['languages']), BODY_X, cv.y, BODY_W)
    cv.skills(load('skills'))
    cv.finish(output)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=ROOT / 'exports' / 'cv-editable.pptx')
    parser.add_argument('--overwrite', action='store_true', help='Replace the output, discarding manual edits to it.')
    args = parser.parse_args()
    if args.output.exists() and not args.overwrite:
        parser.error('Output already exists. Choose --output for a new copy or explicitly pass --overwrite.')
    build(args.output)
